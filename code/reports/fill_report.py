#!/usr/bin/env python3
"""Fill the calibration report from a calibration run.

Reads the config, the markdown report that calibrate.py writes, and (optionally)
a baseline transition, and emits the LaTeX table bodies and figures that
calibration_report.tex \input{}s.  Nothing here re-runs the SMM.

    python fill_report.py --config ../calibration_input_GR.json \
        --calib-report ../output/calibration_growth/reports/calibration_GR_*.md \
        --outdir ../output/calibration_growth --run-baseline --n-sim 1000

Without --run-baseline the figure and the growth table are written with
placeholders, so the report still compiles.
"""
import argparse, json, os, re, sys
import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..'))

NP = r'{\textemdash}'          # placeholder cell


# ---------------------------------------------------------------- parsing ---
def parse_md_table(text, heading):
    """Return the rows of the markdown table under *heading* as lists."""
    m = re.search(rf'^##\s+{re.escape(heading)}\s*$', text, re.M)
    if not m:
        return []
    rest = text[m.end():]
    stop = re.search(r'^##\s+', rest, re.M)
    block = rest[:stop.start()] if stop else rest
    rows = []
    for line in block.strip().split('\n'):
        if not line.startswith('|'):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        if all(set(c) <= set('-: ') for c in cells):      # separator
            continue
        rows.append(cells)
    return rows[1:] if rows else []                        # drop the header


def num(x, default=None):
    try:
        return float(str(x).replace('%', '').replace(',', ''))
    except (TypeError, ValueError):
        return default


def fmt(x, places=4):
    return NP if x is None else f'{x:.{places}f}'


# ----------------------------------------------------------------- tables ---
def params_table(cfg):
    ext, prod, pri, mod = (cfg['external_params'], cfg['production'],
                           cfg['prices'], cfg['model'])
    th = cfg.get('_derived', {}).get('theta', {})
    g, n = ext.get('trend_growth', 0.0), ext.get('pop_growth', 0.0)
    rows_ext = [
        ('$g$', 'labour productivity growth, per capita', g, '2024 Ageing Report'),
        ('$n$', 'population growth', n, 'EUROPOP2023'),
        ('$\\gamma$', 'curvature of consumption utility', mod.get('gamma'), 'log utility'),
        ('$\\varphi$', 'inverse Frisch elasticity', mod.get('phi'), ''),
        ('$\\alpha$', 'private capital share', prod.get('alpha'), ''),
        ('$\\delta$', 'private depreciation', prod.get('delta'), 'standard value'),
        ('$\\eta_g$', 'public capital elasticity', prod.get('eta_g'), ''),
        ('$K_g/Y$', 'public capital ratio', prod.get('K_g'), 'IMF ICSD'),
        ('$\\delta_g$', 'public capital depreciation', prod.get('delta_g'),
         '$I_g/K_g-(\\Gamma-1)$'),
        ('$r$', 'world return on capital', pri.get('r'), ''),
        ('$r_B$', 'sovereign rate', pri.get('r_B'), 'implicit rate 2012--24'),
        ('$\\tau_c,\\tau_l,\\tau_k$', 'consumption, labour, capital tax rates',
         ext.get('tau_c'), 'effective rates'),
        ('$T$, $J_R$', 'lifespan, retirement age', mod.get('T'), ''),
    ]
    rows_cal = [
        ('$\\nu$', 'labour disutility weight', th.get('nu'), 'average hours'),
        ('$\\beta$', 'discount factor', th.get('beta'), '$A/Y$'),
        ('$\\tau_p$', 'payroll tax rate', th.get('tau_p'), 'payroll revenue$/Y$'),
        ('$\\rho^{pens}$', 'pension replacement rate',
         th.get('pension_replacement_default'), 'pensions$/Y$'),
        ('$m^{good}$', 'medical cost scale', th.get('m_good'), 'public health$/Y$'),
    ]
    rows_pin = [
        ('$A$', 'total factor productivity', prod.get('A_tfp'),
         'normalisation, $\\hat y=1$'),
        ('$O/Y$', 'other net spending', cfg['fiscal'].get('other_net_spending_over_Y'),
         'data primary balance'),
    ]
    out = ['\\multicolumn{4}{l}{\\itshape Externally set}\\\\']
    for sym, desc, val, src in rows_ext:
        out.append(f'{sym} & {desc} & {fmt(val)} & {src} \\\\')
    out.append('\\midrule\n\\multicolumn{4}{l}{\\itshape Calibrated jointly by SMM}\\\\')
    for sym, desc, val, src in rows_cal:
        out.append(f'{sym} & {desc} & {fmt(val, 6)} & {src} \\\\')
    out.append('\\midrule\n\\multicolumn{4}{l}{\\itshape Pinned outside the SMM}\\\\')
    for sym, desc, val, src in rows_pin:
        out.append(f'{sym} & {desc} & {fmt(val, 6)} & {src} \\\\')
    return '\n'.join(out)


LABEL = {'average_hours': 'Average hours', 'A_over_Y': '$A/Y$',
         'tax_p_over_Y': 'Payroll revenue$/Y$', 'pensions_over_Y': 'Pensions$/Y$',
         'health_gov_over_Y': 'Public health$/Y$', 'I_g_over_Y': '$I_g/Y$',
         'ui_over_Y': 'UI$/Y$', 'interest_over_Y': 'Interest$/Y$ $(r_BB/Y)$',
         'primary_balance_over_Y': 'Primary balance$/Y$', 'G_over_Y': '$G/Y$',
         'B_over_Y': '$B/Y$', 'health_oop_over_Y': 'Out-of-pocket health$/Y$'}
UNTARGETED = ['I_g_over_Y', 'ui_over_Y', 'interest_over_Y', 'primary_balance_over_Y']


def moments_table(md):
    out = ['\\multicolumn{5}{l}{\\itshape Targeted}\\\\']
    for r in parse_md_table(md, 'Targeted Moments'):
        name, data, model, dev, wt = (r + [''] * 5)[:5]
        out.append(f'{LABEL.get(name, name)} & {fmt(num(data))} & {fmt(num(model))} '
                   f'& {fmt(num(dev), 2)} & {fmt(num(wt), 2)} \\\\')
    ratios = {r[0]: r for r in parse_md_table(md, 'Fiscal Ratios (model vs data, share of Y)')}
    out.append('\\midrule\n\\multicolumn{5}{l}{\\itshape Not targeted}\\\\')
    for key in UNTARGETED:
        r = ratios.get(key)
        if r:
            model, data, dev = num(r[1]), num(r[2]), num(r[3])
        else:
            model = data = dev = None
        out.append(f'{LABEL.get(key, key)} & {fmt(data)} & {fmt(model)} & {fmt(dev, 2)} & \\\\')
    return '\n'.join(out)


def age_table(data_shares, model_t0=None, model_T=None, n_t0=None, n_T=None):
    def col(d, key):
        return NP if d is None else fmt(d.get(key), 3)
    rows = [('Share 25--39', 'young'), ('Share 40--64', 'mid'),
            ('Share 65--84', 'old'), ('Dependency (65--84 / 25--64)', 'oadr')]
    out = []
    for label, key in rows:
        out.append(f'{label} & {col(data_shares, key)} & {col(model_t0, key)} '
                   f'& {col(model_T, key)} \\\\')
    out.append(f'Population growth $n$ & {NP} & {fmt(n_t0, 4)} & {fmt(n_T, 4)} \\\\')
    return '\n'.join(out)


def shares_from_weights(w):
    """w over model ages 0..59 = real 25..84."""
    w = np.asarray(w, float); w = w / w.sum()
    young, mid, old = w[:15].sum(), w[15:40].sum(), w[40:].sum()
    return {'young': young, 'mid': mid, 'old': old, 'oadr': old / (young + mid)}


def growth_table(paths, gamma_minus_1, g, noise=None, t_stable=None):
    """paths: {label: detrended series}.

    The trend is the mean log growth per period over the periods in which the
    population has settled, since only there is the detrended series meant to
    be flat. t_stable is the first such period; without it the second half of
    the horizon is used.
    """
    order = [('$Y$', 'Y'), ('$C$', 'C'), ('$K^{dom}$', 'K_domestic'),
             ('$A$ (wealth)', 'A'), ('$L$', 'L'), ('$K_g$', 'K_g'), ('$B$', 'B')]
    out = []
    for label, key in order:
        x = paths.get(key) if paths else None
        if x is None or len(np.asarray(x)) < 8:
            trend = None
        else:
            x = np.asarray(x, float)
            h = len(x) // 2 if t_stable is None else min(int(t_stable), len(x) - 3)
            trend = 100 * (np.log(x[-1]) - np.log(x[h])) / (len(x) - 1 - h)
        out.append(f'{label} & {fmt(100*gamma_minus_1, 3)} & {fmt(100*g, 3)} '
                   f'& {fmt(trend, 3)} \\\\')
    out.append('\\midrule')
    out.append(f'Theory & {fmt(100*gamma_minus_1, 3)} & {fmt(100*g, 3)} & {{0}} \\\\')
    out.append(f'Noise floor ($g=n=0$) & {{}} & {{}} & {fmt(noise, 3)} \\\\')
    return '\n'.join(out)


def implied_stats(panels, spec, cfg):
    """Model statistics recomputed in the definitions the data use.

    Each calibrated parameter is identified by an aggregate ratio, but the data
    also measure the same object directly, and the two need not agree because
    the bases differ: tau_p applies to the whole wage bill with no contribution
    ceiling, rho_pens multiplies a career-average base rather than the final
    wage. Reporting revenue, base and rate separately shows where the distance
    sits.
    """
    from calibrate import compute_fiscal_ratios
    T = spec.base_config.T
    J_R = spec.base_config.retirement_age
    aw = spec.age_weights if spec.age_weights is not None else np.ones(T) / T
    tot = {k: 0.0 for k in ('wage', 'tax_p', 'pension_new', 'wage_pre',
                            'w_new', 'w_pre')}
    for edu, panel in panels.items():
        sh = spec.education_shares[edu]
        alive = panel.alive_sim.astype(bool)
        for t in range(T):
            a_t = alive[t]
            if not np.any(a_t):
                continue
            wt = sh * aw[t]
            # gross wage income excludes UI, which effective_y_sim includes
            wage = panel.effective_y_sim[t, a_t] - panel.ui_sim[t, a_t]
            tot['wage']  += wt * float(np.mean(wage))
            tot['tax_p'] += wt * float(np.mean(panel.tax_p_sim[t, a_t]))
            if t == J_R:                    # first period of retirement
                tot['pension_new'] += sh * float(np.mean(panel.pension_sim[t, a_t]))
                tot['w_new'] += sh
            if t == J_R - 1:                # last working period
                tot['wage_pre'] += sh * float(np.mean(wage))
                tot['w_pre'] += sh
    fr = compute_fiscal_ratios(panels, spec, cfg)
    Y = None if 'error' in fr else float(fr['Y'])
    ssc = tot['tax_p'] / tot['wage'] if tot['wage'] else None
    base = tot['wage'] / Y if Y else None
    repl = ((tot['pension_new'] / tot['w_new']) / (tot['wage_pre'] / tot['w_pre'])
            if tot['w_new'] and tot['w_pre'] and tot['wage_pre'] else None)
    return {'ssc_rev': None if Y is None else tot['tax_p'] / Y,
            'ssc_base': base, 'ssc_rate': ssc, 'replacement': repl}


# Data counterparts, Greece 2023, as shares. Contribution revenue and the
# factor-income shares are Eurostat (`nasa_10_nf_tr`, `nama_10_gdp`; see
# code/data_inventory.md 1.11); the base is compensation of employees plus
# mixed income, so the self-employed are inside it as they are in the model.
# The replacement rate is the sheet `Parameters` of DATA_GR.xlsx.
DATA_COUNTERPART = [
    ('ssc_rev',     'Social contributions / output',              0.130),
    ('ssc_base',    'Contribution base / output',                 0.570),
    ('ssc_rate',    'Contribution rate on that base',             0.228),
    ('replacement', 'Pension at retirement / last wage',          0.762750),
]


def implied_table(stats, cfg):
    out = []
    for key, label, data in DATA_COUNTERPART:
        model = (stats or {}).get(key)
        dev = (100 * (model / data - 1)) if (model and data) else None
        out.append(f'{label} & {fmt(data, 3)} & {fmt(model, 3)} & {fmt(dev, 1)} \\\\')
    return '\n'.join(out)


# ---------------------------------------------------------------- figures ---
def make_figures(paths, w_data, w_model_t0, outdir):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(1, 2, figsize=(10, 3.6))
    if paths:
        for key, lab in [('Y', r'$\hat y$'), ('C', r'$\hat c$'),
                         ('K_domestic', r'$\hat k^{dom}$'), ('L', r'$\hat\ell$'),
                         ('A', r'$\hat a$')]:
            x = paths.get(key)
            if x is None:
                continue
            x = np.asarray(x, float)
            ax[0].plot(x / x[0], label=lab, lw=1.4)
        ax[0].axhline(1.0, color='0.7', lw=0.8, zorder=0)
        ax[0].set_title('Detrended aggregates, $t=0$ = 1'); ax[0].set_xlabel('period')
        ax[0].legend(frameon=False, ncol=3, fontsize=8)
        Y = np.asarray(paths.get('Y'), float)
        for key, lab in [('B', '$B/Y$'), ('NFA', '$NFA/Y$'),
                         ('K_g', '$K_g/Y$'), ('A', '$A/Y$')]:
            x = paths.get(key)
            if x is None:
                continue
            x = np.asarray(x, float)[:len(Y)]
            ax[1].plot(x / Y[:len(x)], label=lab, lw=1.4)
        ax[1].axhline(0.0, color='0.7', lw=0.8, zorder=0)
        ax[1].set_title('Ratios to output'); ax[1].set_xlabel('period')
        ax[1].legend(frameon=False, ncol=2, fontsize=8)
    for a in ax:
        a.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, 'baseline_panels.pdf'))
    plt.close(fig)

    fig, a = plt.subplots(figsize=(5.2, 3.0))
    ages = np.arange(25, 85)
    if w_data is not None:
        a.plot(ages, np.asarray(w_data) / np.sum(w_data), label='data 2023', lw=1.5)
    if w_model_t0 is not None:
        a.plot(ages, np.asarray(w_model_t0) / np.sum(w_model_t0),
               label='model $t=0$', lw=1.5, ls='--')
    a.set_xlabel('age'); a.set_ylabel('share'); a.legend(frameon=False, fontsize=8)
    a.spines[['top', 'right']].set_visible(False)
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, 'age_distribution.pdf'))
    plt.close(fig)


# ------------------------------------------------------------------- main ---
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--config', default='../calibration_input_GR.json')
    ap.add_argument('--calib-report', default=None,
                    help='markdown report from calibrate.py; newest if omitted')
    ap.add_argument('--outdir', default='../output/calibration_growth')
    ap.add_argument('--run-baseline', action='store_true')
    ap.add_argument('--n-sim', type=int, default=1000)
    ap.add_argument('--backend', default='jax')
    ap.add_argument('--implied', action='store_true',
                    help='compute the data-comparable statistics only')
    args = ap.parse_args()

    cfg = json.load(open(args.config))
    outdir = os.path.abspath(args.outdir)
    os.makedirs(outdir, exist_ok=True)

    report = args.calib_report
    if report is None:
        d = os.path.join(outdir, 'reports')
        cands = sorted(f for f in os.listdir(d) if f.endswith('.md')) if os.path.isdir(d) else []
        report = os.path.join(d, cands[-1]) if cands else None
    md = open(report).read() if report else ''
    if report:
        print(f'calibration report: {os.path.basename(report)}')

    g = cfg['external_params'].get('trend_growth', 0.0)
    n = cfg['external_params'].get('pop_growth', 0.0)
    gamma_minus_1 = (1 + g) * (1 + n) - 1

    # First period on the balanced growth path, from the demographic sidecar.
    t_stable = None
    demog = os.path.join(os.path.dirname(args.config), '..', 'data',
                         'demography_GR.npz')
    if os.path.exists(demog):
        _d = np.load(demog)
        t_stable = int(_d['stable_year']) - int(_d['base_year'])

    # data age distribution (25-84), cached next to the data if openpyxl is absent
    w_data = None
    cache = os.path.join(os.path.dirname(args.config), '..', 'data', 'agedist_2023.npy')
    if os.path.exists(cache):
        w_data = np.load(cache)
    else:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(os.path.join(os.path.dirname(args.config), '..',
                                                     'data', 'DATA_GR.xlsx'),
                                        read_only=True, data_only=True)
            ws = wb['Population by age']; rows = list(ws.iter_rows(values_only=True))
            hdr = [str(c) for c in rows[8]]
            def colof(a):
                lbl = 'Less than 1 year' if a == 0 else ('1 year' if a == 1 else f'{a} years')
                for j, h in enumerate(hdr):
                    if h.strip() == lbl:
                        return j
            row23 = next(r for r in rows[10:]
                         if r[1] is not None and str(r[1])[:4] == '2023')
            w_data = np.array([row23[colof(a)] if isinstance(row23[colof(a)], (int, float))
                               else 0.0 for a in range(25, 85)], float)
            np.save(cache, w_data / w_data.sum())
        except Exception as e:                                   # noqa: BLE001
            print(f'  age distribution unavailable: {e}')

    # The age weights come from the demographic path, so the age-structure
    # table and its figure need no solve; only the panels and the growth table
    # need a baseline transition.
    from calibrate import load_config, build_olg_transition
    L = load_config(args.config)
    economy, tp, T_TR = build_olg_transition(L['config_data'], backend=args.backend)

    def living_shares(t):
        """Shares of the living population by age at period t.

        The transition's weights are cohort sizes at entry, because its means
        run over all agents with the dead holding zero. The data column counts
        the living, so the weights are carried through cumulative survival
        before the two are compared.
        """
        w = economy._cohort_weights(t)
        cum = np.empty(economy.T)
        for j in range(economy.T):
            sched = economy._cohort_survival_schedule(t - j)
            cum[j] = float(np.prod(np.mean(sched, axis=1)[:j])) if j else 1.0
        out = w * cum
        return out / out.sum()

    w_model_t0 = living_shares(0)
    w_model_T = living_shares(T_TR - 1)
    G_path = economy.growth_factors(T_TR)
    n = float(G_path[0] / (1.0 + g) - 1.0)          # realised, not the config scalar
    n_T = float(G_path[-1] / (1.0 + g) - 1.0)

    paths = None
    if args.run_baseline:
        prod = L['config_data']['production']
        I_g = ((prod.get('delta_g', 0.05) + economy.growth_factors(T_TR) - 1.0)
               * prod.get('K_g', 0.0))
        tax = {k: tp[k] for k in ('tau_c_path', 'tau_l_path', 'tau_p_path',
                                  'tau_k_path', 'pension_replacement_path')}
        print(f'baseline transition: T={T_TR}, n_sim={args.n_sim}, backend={args.backend}')
        res = economy.simulate_transition(r_path=tp['r_path'], I_g_path=I_g,
                                          n_sim=args.n_sim, verbose=False, **tax)
        paths = {k: np.asarray(v) for k, v in res.items()
                 if isinstance(v, (list, np.ndarray)) and np.ndim(v) == 1}

    stats = None
    if args.run_baseline or args.implied:
        from calibrate import run_model_moments
        import dataclasses
        L2 = L
        sp = dataclasses.replace(L2['spec'], backend=args.backend, n_sim=args.n_sim)
        th = np.array([cfg['_derived']['theta'][p.name] for p in sp.params])
        print(f'implied statistics: solving at n_sim={args.n_sim} ...')
        _, panels = run_model_moments(th, sp, return_panels=True)
        # load_config injects _derived.K_over_L; the raw JSON has it as None,
        # which makes compute_fiscal_ratios return an error instead of Y.
        stats = implied_stats(panels, sp, L2['config_data'])
        for key, lbl, dat in DATA_COUNTERPART:
            v = stats.get(key)
            print(f'  {lbl:40s} model {v if v is not None else float("nan"):.4f}'
                  f'   data {dat:.4f}')

    def wrap(colspec, header, body):
        return ('\\begin{tabular}{@{}' + colspec + '@{}}\n\\toprule\n'
                + header + ' \\\\\n\\midrule\n' + body
                + '\n\\bottomrule\n\\end{tabular}')

    frag = {
        'implied_body.tex': wrap(
            'lS[table-format=1.3]S[table-format=1.3]S[table-format=+3.1]',
            ' & {Data} & {Model} & {\\% dev}', implied_table(stats, cfg)),
        'params_body.tex': wrap(
            'llS[table-format=2.4]l',
            'Symbol & Description & {Value} & Source / identified by',
            params_table(cfg)),
        'moments_body.tex': wrap(
            'lS[table-format=1.4]S[table-format=1.4]S[table-format=+2.2]S[table-format=3.2]',
            ' & {Data} & {Model} & {\\% dev} & {Weight}', moments_table(md)),
        'age_body.tex': wrap(
            'lS[table-format=1.3]S[table-format=1.3]S[table-format=1.3]',
            ' & {Data, 2023} & {Model, $t=0$} & {Model, terminal}',
            age_table(shares_from_weights(w_data) if w_data is not None else None,
                      shares_from_weights(w_model_t0) if w_model_t0 is not None else None,
                      shares_from_weights(w_model_T) if w_model_T is not None else None,
                      n, n_T)),
        'growth_body.tex': wrap(
            'lS[table-format=1.3]S[table-format=1.3]S[table-format=+1.3]',
            ' & {Level (\\%)} & {Per capita (\\%)} & {Detrended trend (\\%)}',
            growth_table(paths, gamma_minus_1, g, t_stable=t_stable)),
    }
    for name, body in frag.items():
        with open(os.path.join(outdir, name), 'w') as fh:
            fh.write(body + '\n')
        print(f'  wrote {name}')

    make_figures(paths, w_data, w_model_t0, outdir)
    print('  wrote baseline_panels.pdf, age_distribution.pdf')

    meta = {'run': os.path.basename(report) if report else None,
            'g': g, 'n': n, 'Gamma_minus_1': gamma_minus_1,
            'r_B': cfg['prices'].get('r_B'),
            'objective': num((parse_md_table(md, 'Summary') or [['', '']])[0][1])}
    with open(os.path.join(outdir, 'header.tex'), 'w') as fh:
        run = (meta['run'] or '').replace('_', '\\_') or NP
        for name, val in (('runlabel', run), ('gval', fmt(g, 4)),
                          ('nval', fmt(n, 4)), ('rBval', fmt(meta['r_B'], 4)),
                          ('Gammaval', fmt(gamma_minus_1, 5))):
            fh.write(f"\\providecommand{{\\{name}}}{{}}"
                     f"\\renewcommand{{\\{name}}}{{{val}}}\n")
    print('  wrote header.tex')


if __name__ == '__main__':
    main()
